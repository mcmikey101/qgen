from io import BytesIO
from pathlib import Path

from fastapi import HTTPException

MAX_CHUNK_CHARS = 12000
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}


def extract_text(filename: str, contents: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise HTTPException(status_code=415, detail=f"Unsupported file type. Use: {allowed}")

    try:
        if extension == ".pdf":
            from pypdf import PdfReader

            extracted = "\n\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(contents)).pages)
        elif extension == ".docx":
            from docx import Document

            document = Document(BytesIO(contents))
            extracted = "\n\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                extracted += "\n\n" + "\n".join(
                    "\t".join(cell.text for cell in row.cells) for row in table.rows
                )
        else:
            extracted = contents.decode("utf-8-sig")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Could not read the uploaded file") from exc

    text = "\n".join(line.rstrip() for line in extracted.splitlines()).strip()
    if not text:
        raise HTTPException(status_code=422, detail="The file contains no readable text")
    return text


def split_text(text: str, max_chars: int = MAX_CHUNK_CHARS) -> list[str]:
    """Split on paragraph boundaries, with hard splits for oversized paragraphs."""
    chunks: list[str] = []
    current = ""

    for paragraph in (part.strip() for part in text.split("\n\n")):
        if not paragraph:
            continue
        pieces = [paragraph[i : i + max_chars] for i in range(0, len(paragraph), max_chars)]
        for piece in pieces:
            candidate = f"{current}\n\n{piece}" if current else piece
            if len(candidate) > max_chars:
                chunks.append(current)
                current = piece
            else:
                current = candidate

    if current:
        chunks.append(current)
    return chunks
